# -*- coding: utf-8 -*-
"""
生成 Word 接口测试报告（严格按照模板结构）：
一、测试概要（2 列 key-value 表）
二、接口用例统计（5 列：接口名称/请求路径/用例数/通过/失败）
三、失败用例及问题分析（按接口分组的二级小节，每小节一张 5 列表）
四、修复建议（按接口分组 bullet list）
五、测试覆盖说明（2 列表格：测试类型/说明）
六、附录（测试用例分布 + 对应 Excel 文件名）
"""
import sys
import io
import json
import argparse
from pathlib import Path
from datetime import datetime
from collections import OrderedDict, Counter

sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8', errors='replace')

from docx import Document
from docx.shared import Pt, RGBColor, Cm, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH

SKILL_DIR = Path(__file__).resolve().parent.parent
RESULTS_DIR = SKILL_DIR / "test_results"
REPORTS_DIR = SKILL_DIR / "reports"


def classify_failure(http, expected, actual_msg, case_type):
    """按模板给定 5 类失败分类之外，这里根据响应再细分为更可读的问题描述。"""
    msg = (actual_msg or "").lower()
    expected_l = (expected or "").lower()

    if http == 0:
        return "网络问题"
    # 预期应返回 401/500 等，实际返回 200 —— 鉴权/参数校验缺失
    if "token" in expected_l or "mobile" in expected_l or "401" in expected_l:
        return "鉴权校验缺失"
    if "500" in expected_l or "提示参数" in expected or "必填" in expected or "缺少" in expected:
        if http == 200:
            return "参数校验缺失"
    if "超长" in expected or "超出" in expected or "长度" in expected:
        return "参数校验缺失"
    if "格式" in expected and http == 200:
        return "缺少参数格式校验"
    if "比例" in expected or "金额" in expected:
        return "参数校验缺失"
    if "不存在" in expected and http == 200:
        return "未校验数据存在性"
    if "拦截" in expected or "SQL" in (actual_msg or "").upper():
        return "安全校验待加强"
    if "不可删除" in expected:
        return "业务规则校验缺失"
    return "业务校验缺失"


def set_cell_style(cell, *, bold=False, color=None, size=10, align=None):
    for p in cell.paragraphs:
        if align is not None:
            p.alignment = align
        for run in p.runs:
            run.font.size = Pt(size)
            run.font.name = "微软雅黑"
            if bold:
                run.bold = True
            if color:
                run.font.color.rgb = RGBColor.from_string(color)


def add_kv_table(doc, rows):
    tbl = doc.add_table(rows=len(rows), cols=2)
    tbl.style = "Table Grid"
    for i, (k, v) in enumerate(rows):
        c0 = tbl.rows[i].cells[0]
        c1 = tbl.rows[i].cells[1]
        c0.text = ""
        c1.text = ""
        c0.paragraphs[0].add_run(str(k)).bold = True
        c1.paragraphs[0].add_run(str(v))
        set_cell_style(c0, bold=True, size=10.5)
        set_cell_style(c1, size=10.5)
        # Width
        c0.width = Cm(4)
        c1.width = Cm(12)
    return tbl


def add_header_table(doc, headers, rows, widths=None):
    tbl = doc.add_table(rows=1 + len(rows), cols=len(headers))
    tbl.style = "Table Grid"
    hdr_cells = tbl.rows[0].cells
    for i, h in enumerate(headers):
        hdr_cells[i].text = ""
        run = hdr_cells[i].paragraphs[0].add_run(str(h))
        run.bold = True
        run.font.size = Pt(10.5)
        run.font.name = "微软雅黑"
        hdr_cells[i].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        # Background color via shading element
        from docx.oxml.ns import qn
        from docx.oxml.shared import OxmlElement
        tc_pr = hdr_cells[i]._tc.get_or_add_tcPr()
        shd = OxmlElement('w:shd')
        shd.set(qn('w:val'), 'clear')
        shd.set(qn('w:color'), 'auto')
        shd.set(qn('w:fill'), '4472C4')
        tc_pr.append(shd)
        # White font
        for p in hdr_cells[i].paragraphs:
            for r in p.runs:
                r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    for ri, row in enumerate(rows, start=1):
        for ci, val in enumerate(row):
            c = tbl.rows[ri].cells[ci]
            c.text = ""
            run = c.paragraphs[0].add_run(str(val) if val is not None else "")
            run.font.size = Pt(10)
            run.font.name = "微软雅黑"
            if ci in (0, len(headers) - 1) or len(str(val)) < 20:
                c.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    if widths:
        for ci, w in enumerate(widths):
            for row in tbl.rows:
                row.cells[ci].width = Cm(w)
    return tbl


def generate(results_file: Path, output_file: Path):
    with open(results_file, "r", encoding="utf-8") as f:
        data = json.load(f)

    project_name = data.get("project", "接口测试")
    base_url = data.get("base_url", "")
    end_time = data.get("end_time") or datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    cases = data["cases"]

    # 分组：单接口按 api_name 聚合，业务流程用例单独统计
    single_apis = OrderedDict()   # api_name -> list[case]
    flow_cases = []              # 业务流程用例列表（case_type == "业务流程"）
    for c in cases:
        if c.get("case_type", "") == "业务流程":
            flow_cases.append(c)
        else:
            single_apis.setdefault(c["api_name"], []).append(c)

    total = len(cases)
    passed = sum(1 for c in cases if c["result"] == "PASS")
    failed = total - passed

    # 计算对应 Excel 文件名
    dt = datetime.strptime(end_time, "%Y-%m-%d %H:%M:%S")
    ts = dt.strftime("%Y%m%d_%H%M%S")
    excel_name = f"{project_name}接口测试用例及测试结果_{ts}.xlsx"

    # 开始写文档
    doc = Document()

    # 标题
    title = doc.add_heading(f"{project_name}接口测试报告", level=0)
    for run in title.runs:
        run.font.name = "微软雅黑"

    # 一、测试概要
    doc.add_heading("一、测试概要", level=1)
    api_count_text = f"{len(single_apis)}个（{'/'.join(single_apis.keys())}）"
    summary_rows = [
        ("测试时间", end_time),
        ("服务地址", base_url),
        ("接口总数", api_count_text),
        ("业务流程数", f"{len(flow_cases)}条"),
        ("用例总数", f"{total}条"),
        ("通过数量", f"{passed}条"),
        ("失败数量", f"{failed}条"),
        ("通过率", f"{passed/total*100:.1f}%"),
        ("测试类型", "正向测试、逆向测试、边界测试、安全测试、业务流程测试"),
    ]
    add_kv_table(doc, summary_rows)
    doc.add_paragraph("")

    # 二、接口用例统计
    doc.add_heading("二、接口用例统计", level=1)
    stat_rows = []
    for name, lst in single_apis.items():
        path = lst[0]["path"]
        p_cnt = sum(1 for c in lst if c["result"] == "PASS")
        f_cnt = sum(1 for c in lst if c["result"] == "FAIL")
        stat_rows.append([name, path, str(len(lst)), str(p_cnt), str(f_cnt)])
    # 业务流程合并为一行
    if flow_cases:
        p_cnt = sum(1 for c in flow_cases if c["result"] == "PASS")
        f_cnt = sum(1 for c in flow_cases if c["result"] == "FAIL")
        stat_rows.append(["业务流程（多接口串联）", "(多步骤串联）", str(len(flow_cases)), str(p_cnt), str(f_cnt)])
    add_header_table(doc, ["接口/流程名称", "请求路径", "用例数", "通过", "失败"],
                     stat_rows, widths=[5, 5.5, 2, 2, 2])
    doc.add_paragraph("")

    # 三、失败用例及问题分析
    doc.add_heading("三、失败用例及问题分析", level=1)

    # 单接口
    api_idx = 0
    for name, lst in single_apis.items():
        fails = [c for c in lst if c["result"] == "FAIL"]
        api_idx += 1
        doc.add_heading(f"3.{api_idx} {name}", level=2)
        if not fails:
            doc.add_paragraph("本接口无失败用例。")
            continue
        fail_rows = []
        for c in fails:
            reason = classify_failure(c["http_status"],
                                      c["expected"], c["actual_msg"], c["case_type"])
            actual = f"HTTP {c['http_status']}, {c.get('actual_msg', '')}"
            fail_rows.append([c["case_id"], c["desc"], c["expected"], actual, reason])
        add_header_table(doc, ["用例编号", "测试描述", "预期结果", "实际结果", "问题分析"],
                         fail_rows, widths=[2, 5, 5, 3, 3])

    # 业务流程失败用例
    flow_fails = [c for c in flow_cases if c["result"] == "FAIL"]
    if flow_fails:
        api_idx += 1
        doc.add_heading(f"3.{api_idx} 业务流程（多接口串联）", level=2)
        fail_rows = []
        for c in flow_fails:
            reason = classify_failure(c["http_status"],
                                      c["expected"], c["actual_msg"], c["case_type"])
            actual = f"HTTP {c['http_status']}, {c.get('actual_msg', '')}"
            fail_rows.append([c["case_id"], c["desc"], c["expected"], actual, reason])
        add_header_table(doc, ["用例编号", "测试描述", "预期结果", "实际结果", "问题分析"],
                         fail_rows, widths=[2, 5, 5, 3, 3])

    # 四、修复建议
    doc.add_heading("四、修复建议", level=1)

    def add_suggest_group(title_name, fails):
        if not fails:
            return
        doc.add_paragraph(f"【{title_name}】")
        # 聚类按 reason
        counter = Counter()
        for c in fails:
            reason = classify_failure(c["http_status"],
                                      c["expected"], c["actual_msg"], c["case_type"])
            counter[reason] += 1
        for reason, cnt in counter.most_common():
            p = doc.add_paragraph(style="List Bullet")
            p.add_run(f"• {reason}（{cnt}条用例）")

    for name, lst in single_apis.items():
        fails = [c for c in lst if c["result"] == "FAIL"]
        add_suggest_group(name, fails)
    if flow_cases:
        fails = [c for c in flow_cases if c["result"] == "FAIL"]
        add_suggest_group("业务流程（多接口串联）", fails)

    doc.add_paragraph("")

    # 五、测试覆盖说明
    doc.add_heading("五、测试覆盖说明", level=1)
    cover_rows = [
        ["正向测试", "验证正常业务流程，包括完整合法数据、默认参数、不同组合（合同类型/等级/签约主体等）"],
        ["逆向测试", "验证异常场景处理，包括必填字段缺失、字段为空、不存在的 uuid/fileId、非法参数值等"],
        ["边界测试", "验证边界值处理，包括最小/最大长度、超长字符、空字符串、特殊字符、超大分页、非法日期格式等"],
        ["安全测试", "验证安全性，包括 SQL 注入、XSS 脚本攻击等"],
        ["业务流程测试", "验证跨接口端到端流程，覆盖：草稿创建→查询→详情→编辑→开启流程→删除限制 等完整业务链路"],
    ]
    add_header_table(doc, ["测试类型", "说明"], cover_rows, widths=[3, 13])
    doc.add_paragraph("")

    # 六、附录
    doc.add_heading("六、附录", level=1)
    doc.add_paragraph("测试用例分布：")
    for name, lst in single_apis.items():
        type_counts = Counter(c["case_type"] for c in lst)
        parts = [f"{t}{cnt}条" for t, cnt in sorted(type_counts.items())]
        p = doc.add_paragraph(style="List Bullet")
        p.add_run(f"• {name}：{len(lst)}条（{'+'.join(parts)}）")
    if flow_cases:
        type_counts = Counter(c["case_type"] for c in flow_cases)
        parts = [f"{t}{cnt}条" for t, cnt in sorted(type_counts.items())]
        p = doc.add_paragraph(style="List Bullet")
        p.add_run(f"• 业务流程（多接口串联）：{len(flow_cases)}条（{'+'.join(parts)}）")
    p = doc.add_paragraph(style="List Bullet")
    p.add_run(f"• 总计：{total}条测试用例")
    doc.add_paragraph("")
    doc.add_paragraph(f"详细测试用例请查看：{excel_name}")

    doc.save(output_file)
    print(f"[OK] Word 报告已生成: {output_file}")


def main():
    parser = argparse.ArgumentParser(description="Generate Word test report")
    parser.add_argument("--workspace", dest="workspace", default=None,
                        help="Path to workspace directory (overrides default paths)")
    parser.add_argument("--output", dest="output_dir", default=None,
                        help="Output directory for reports (default: same as workspace)")
    parser.add_argument("--results-file", dest="results_file", default=None,
                        help="Path to results JSON (default: workspace/test_results/latest_results.json)")
    args = parser.parse_args()

    # Resolve results and reports dirs from workspace or defaults
    if args.workspace:
        ws = Path(args.workspace)
    else:
        ws_root = SKILL_DIR / "workspaces"
        try:
            subdirs = [d for d in ws_root.iterdir() if d.is_dir() and not d.name.startswith(('.', '_'))]
            ws = subdirs[0] if len(subdirs) == 1 else None
        except FileNotFoundError:
            ws = None

    # 如果指定了 --output，直接使用 output 目录
    if args.output_dir:
        ws = Path(args.output_dir)

    if ws:
        results_dir = ws / "test_results"
        reports_dir = ws / "reports"
    else:
        results_dir = RESULTS_DIR
        reports_dir = REPORTS_DIR
    reports_dir.mkdir(parents=True, exist_ok=True)

    if args.results_file:
        latest = Path(args.results_file)
    else:
        latest = results_dir / "latest_results.json"

    if not latest.exists():
        print("[ERR] 找不到 latest_results.json，请先运行测试脚本")
        sys.exit(1)
    with open(latest, "r", encoding="utf-8") as f:
        data = json.load(f)
    project = data.get("project", "接口测试")
    end_time = data.get("end_time") or datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    dt = datetime.strptime(end_time, "%Y-%m-%d %H:%M:%S")
    ts = dt.strftime("%Y%m%d_%H%M%S")
    out = reports_dir / f"{project}接口测试报告_{ts}.docx"
    generate(latest, out)


if __name__ == "__main__":
    main()
