# -*- coding: utf-8 -*-
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from lxml import etree

def set_cell_shading(cell, color):
    """设置单元格背景色"""
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    # 创建shading元素
    shading = etree.Element(qn('w:shd'))
    shading.set(qn('w:fill'), color)
    shading.set(qn('w:val'), 'clear')
    tcPr.append(shading)

def set_cell_text(cell, text, bold=False, size=9, color=None, align=WD_ALIGN_PARAGRAPH.LEFT):
    """设置单元格文本"""
    cell.text = ''
    p = cell.paragraphs[0]
    p.alignment = align
    run = p.add_run(text)
    run.bold = bold
    run.font.size = Pt(size)
    if color:
        run.font.color.rgb = RGBColor(*color)

def create_table_with_style(doc, headers, rows, col_widths, header_color='1F4E79'):
    """创建格式化表格"""
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = 'Table Grid'
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    
    # 设置列宽
    for i, width in enumerate(col_widths):
        for row in table.rows:
            row.cells[i].width = Cm(width)
    
    # 表头
    header_row = table.rows[0]
    for i, header in enumerate(headers):
        cell = header_row.cells[i]
        set_cell_text(cell, header, bold=True, size=10, color=(255,255,255), align=WD_ALIGN_PARAGRAPH.CENTER)
        set_cell_shading(cell, header_color)
    
    # 数据行
    for row_idx, row_data in enumerate(rows):
        row = table.rows[row_idx + 1]
        for col_idx, cell_text in enumerate(row_data):
            cell = row.cells[col_idx]
            set_cell_text(cell, str(cell_text), size=9)
    
    return table

# 创建文档
doc = Document()

# 设置默认字体
style = doc.styles['Normal']
style.font.name = '仿宋'
style._element.rPr.rFonts.set(qn('w:eastAsia'), '仿宋')
style.font.size = Pt(12)

# ===== 标题 =====
title = doc.add_paragraph()
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = title.add_run('需求文档质量评估报告')
run.bold = True
run.font.size = Pt(22)
run.font.name = '黑体'
run._element.rPr.rFonts.set(qn('w:eastAsia'), '黑体')

# ===== 基本信息 =====
doc.add_paragraph()
info = [
    ('文档名称', '广东省交通基础设施数字化转型管理系统需求分析说明书'),
    ('评估时间', '2026-04-14'),
    ('综合评分', '57/100（等级：D级 - 较差）'),
]
for label, value in info:
    p = doc.add_paragraph()
    run = p.add_run(f'{label}：')
    run.bold = True
    run.font.size = Pt(12)
    p.add_run(value).font.size = Pt(12)

doc.add_paragraph('_' * 80)

# ===== 一、维度评分 =====
h1 = doc.add_heading('一、维度评分', level=1)
h1.runs[0].font.name = '黑体'
h1.runs[0]._element.rPr.rFonts.set(qn('w:eastAsia'), '黑体')

headers = ['维度', '得分', '权重', '加权分', '说明']
rows = [
    ['可测试性', '3/5', '30%', '18', '部分满足，非功能需求缺失严重'],
    ['完整性', '2/5', '25%', '10', '严重不足，影响测试设计'],
    ['清晰性', '3/5', '20%', '12', '存在多处模糊表述'],
    ['一致性', '4/5', '10%', '8', '基本一致，轻微表述差异'],
    ['可追溯性', '3/5', '10%', '6', '有编号，缺版本追溯'],
    ['可实现性', '3/5', '5%', '3', '存在技术风险待评估'],
    ['综合得分', '57/100', '100%', '57', 'D级 - 较差'],
]
table = create_table_with_style(doc, headers, rows, [2.5, 1.5, 1.5, 1.5, 5], '1F4E79')

# 标记最后一行
last_row = table.rows[-1]
for cell in last_row.cells:
    set_cell_shading(cell, 'E0E0E0')
    if cell.paragraphs[0].runs:
        cell.paragraphs[0].runs[0].bold = True

doc.add_paragraph('_' * 80)

# ===== 二、问题清单 =====
h1 = doc.add_heading('二、问题清单', level=1)
h1.runs[0].font.name = '黑体'
h1.runs[0]._element.rPr.rFonts.set(qn('w:eastAsia'), '黑体')

# 2.1 阻塞问题
h2 = doc.add_heading('2.1 阻塞问题（无法进行测试设计）', level=2)
h2.runs[0].font.name = '黑体'
h2.runs[0]._element.rPr.rFonts.set(qn('w:eastAsia'), '黑体')

headers = ['序号', '维度', '问题描述', '详细问题', '建议']
rows = [
    ['1', '完整性', '非功能需求缺失', '用户界面、性能、安全、软硬件环境、兼容性需求章节全部为空', '补充具体性能指标、安全要求等'],
    ['2', '可测试性', '性能指标缺失', '无法设计性能测试用例：响应时间、并发数均未定义', '明确页面加载<3秒、数据更新<10秒等指标'],
    ['3', '完整性', '异常场景缺失', '接口异常、数据为空、网络超时等处理逻辑缺失', '补充异常处理流程和重试机制'],
    ['4', '安全需求', '安全需求为空', '认证方式、会话管理、权限控制、数据加密均未说明', '补充密码规则、登录锁定策略、会话超时等'],
]
table = create_table_with_style(doc, headers, rows, [1, 1.8, 2, 5, 3.5], 'C00000')
for row in table.rows[1:]:
    for cell in row.cells:
        set_cell_shading(cell, 'FFE6E6')

doc.add_paragraph()

# 2.2 警告问题
h2 = doc.add_heading('2.2 警告问题（影响测试覆盖率）', level=2)
h2.runs[0].font.name = '黑体'
h2.runs[0]._element.rPr.rFonts.set(qn('w:eastAsia'), '黑体')

headers = ['序号', '维度', '问题描述', '详细问题', '建议']
rows = [
    ['5', '可测试性', '登录规则不完整', '仅提到密码长度>=8位，缺少格式、复杂度、锁定规则', '补充用户名格式、密码复杂度、连续失败锁定'],
    ['6', '完整性', '接口规范缺失', '缺少接口地址、请求方式、参数格式、返回格式、错误码', '补充接口文档：URL、参数、返回、错误码'],
    ['7', '清晰性', '数据更新规则模糊', '"严格按用户需求执行更新"含义不明确', '改为明确规则：刷新间隔不低于5分钟'],
    ['8', '可测试性', '边界条件未定义', '任务数量上限、数据量上限、地图点位上限均未说明', '补充数量上限定义'],
    ['9', '完整性', '并发场景缺失', '多用户同时操作、数据同时更新等场景处理逻辑缺失', '补充并发控制、缓存机制、锁策略'],
    ['10', '可追溯性', '修订记录为空', '文档修订记录表为空，无法了解需求变更历史', '填写版本号、修订日期、修订内容'],
    ['11', '完整性', '备份恢复缺失', '数据备份恢复需求缺失', '补充备份频率、保留周期、恢复流程'],
    ['12', '可测试性', '空数据处理缺失', '里程为0，投资为0时页面如何展示未定义', '补充空数据展示规则'],
]
table = create_table_with_style(doc, headers, rows, [1, 1.8, 2, 5, 3.5], 'FF8C00')
for row in table.rows[1:]:
    for cell in row.cells:
        set_cell_shading(cell, 'FFF3CD')

doc.add_paragraph()

# 2.3 建议问题
h2 = doc.add_heading('2.3 建议问题（可优化项）', level=2)
h2.runs[0].font.name = '黑体'
h2.runs[0]._element.rPr.rFonts.set(qn('w:eastAsia'), '黑体')

headers = ['序号', '维度', '问题描述', '详细问题', '建议']
rows = [
    ['13', '清晰性', '术语解释表为空', '术语缩写解释表（表1-2）为空', '补充"数转"、"部绩效系统"等缩写'],
    ['14', '清晰性', '参考文献缺失', '参考文献表（表1-1）为空', '补充相关通知文件、考核指标文件'],
    ['15', '一致性', '编号格式不一致', '功能编号偶有不一致', '统一使用完整编号格式'],
    ['16', '可追溯性', '需求来源追溯缺失', '未关联到具体业务需求或用户故事', '增加"需求来源"字段'],
]
table = create_table_with_style(doc, headers, rows, [1, 1.8, 2, 5, 3.5], '4472C4')
for row in table.rows[1:]:
    for cell in row.cells:
        set_cell_shading(cell, 'E7F3FF')

doc.add_paragraph('_' * 80)

# ===== 三、改进建议 =====
h1 = doc.add_heading('三、改进建议', level=1)
h1.runs[0].font.name = '黑体'
h1.runs[0]._element.rPr.rFonts.set(qn('w:eastAsia'), '黑体')

# P0
h2 = doc.add_heading('P0 - 必须立即补充（阻塞测试设计）', level=2)
h2.runs[0].font.name = '黑体'
h2.runs[0]._element.rPr.rFonts.set(qn('w:eastAsia'), '黑体')
p0_items = [
    '补充性能需求：页面响应时间、数据刷新耗时、并发用户数',
    '补充安全需求：认证方式、密码规则、会话管理、权限控制',
    '补充异常处理：接口超时、数据为空、网络异常的处理逻辑',
    '补充接口规范：数字底座接口的详细参数、返回格式、错误码',
]
for item in p0_items:
    p = doc.add_paragraph(style='List Bullet')
    p.add_run(item).font.size = Pt(11)

# P1
doc.add_paragraph()
h2 = doc.add_heading('P1 - 应尽快补充（影响测试覆盖率）', level=2)
h2.runs[0].font.name = '黑体'
h2.runs[0]._element.rPr.rFonts.set(qn('w:eastAsia'), '黑体')
p1_items = [
    '补充边界条件定义',
    '补充并发控制策略',
    '补充数据备份恢复需求',
    '补充日志审计需求',
]
for item in p1_items:
    p = doc.add_paragraph(style='List Bullet')
    p.add_run(item).font.size = Pt(11)

# P2
doc.add_paragraph()
h2 = doc.add_heading('P2 - 建议优化', level=2)
h2.runs[0].font.name = '黑体'
h2.runs[0]._element.rPr.rFonts.set(qn('w:eastAsia'), '黑体')
p2_items = [
    '填写术语解释表',
    '填写参考文献',
    '填写文档修订记录',
    '补充需求来源追溯',
]
for item in p2_items:
    p = doc.add_paragraph(style='List Bullet')
    p.add_run(item).font.size = Pt(11)

doc.add_paragraph('_' * 80)

# ===== 四、评估等级说明 =====
h1 = doc.add_heading('四、评估等级说明', level=1)
h1.runs[0].font.name = '黑体'
h1.runs[0]._element.rPr.rFonts.set(qn('w:eastAsia'), '黑体')

headers = ['分数段', '等级', '说明', '建议']
rows = [
    ['90-100', 'A（优秀）', '可直接进入测试设计', '-'],
    ['75-89', 'B（良好）', '小幅修订后可用', '针对性改进'],
    ['60-74', 'C（一般）', '需补充完善后使用', '重点补充缺失项'],
    ['40-59', 'D（较差）', '需重大修订', '重点解决阻塞问题'],
    ['<40', 'E（不合格）', '建议重写', '重新梳理需求'],
]
table = create_table_with_style(doc, headers, rows, [2, 2.5, 3.5, 3], '1F4E79')

# 标记D级行
d_row = table.rows[4]
for cell in d_row.cells:
    set_cell_shading(cell, 'FFF3CD')

# 标记E级行
e_row = table.rows[5]
for cell in e_row.cells:
    set_cell_shading(cell, 'FFE6E6')

doc.add_paragraph('_' * 80)

# ===== 五、评估结论 =====
h1 = doc.add_heading('五、评估结论', level=1)
h1.runs[0].font.name = '黑体'
h1.runs[0]._element.rPr.rFonts.set(qn('w:eastAsia'), '黑体')

p = doc.add_paragraph()
p.add_run('当前需求文档非功能需求严重缺失，建议先补充P0级别内容后再开展测试设计工作。').font.size = Pt(12)

doc.add_paragraph()
doc.add_paragraph()

# 签字栏
p = doc.add_paragraph()
p.add_run('评估人：________________    日期：________________').font.size = Pt(12)

# 保存文档
import os
script_dir = os.path.dirname(os.path.abspath(__file__))
skill_dir = os.path.dirname(script_dir)
output_dir = os.path.join(skill_dir, 'analysis_report')
os.makedirs(output_dir, exist_ok=True)
import datetime
date_str = datetime.date.today().strftime('%Y%m%d')
output_path = os.path.join(output_dir, f'广东省交通基础设施数字化转型系统需求评估报告_{date_str}.docx')
doc.save(output_path)
print(f'文档生成成功：{output_path}')
