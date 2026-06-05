# -*- coding: utf-8 -*-
"""
{模块名}测试脚本
自动生成模板，请根据实际项目修改配置区和用例分支
"""
import os
import time
import shutil
from datetime import datetime
from playwright.sync_api import sync_playwright
from openpyxl import load_workbook
from openpyxl.styles import PatternFill
from openpyxl.drawing.image import Image as XLImage
from openpyxl.drawing.spreadsheet_drawing import OneCellAnchor

# ==================== 配置区（每个项目需修改） ====================
CDP_URL = "http://127.0.0.1:9222"
TARGET_URL = "{页面URL}"
USERNAME = "{账号}"
PASSWORD = "{密码}"
BASE_INPUT = r"{项目根目录}\input"
BASE_OUTPUT = r"{项目根目录}\results"
BASE_SCREENSHOTS = r"{项目根目录}\screenshots"
CASE_FILE = "{测试用例Excel文件名}.xlsx"
START_ROW = 10  # Excel中用例起始行
# ==================== 配置区结束 ====================


# ==================== 辅助函数 ====================
def safe_filename(name):
    invalid = '/\\:*?"<>|'
    for c in invalid:
        name = name.replace(c, '_')
    return name


def load_cases():
    """从Excel加载测试用例"""
    wb = load_workbook(os.path.join(BASE_INPUT, CASE_FILE), data_only=True)
    ws = wb.active
    cases = []
    for row_idx in range(START_ROW, ws.max_row + 1):
        case_id = ws.cell(row_idx, 3).value
        if case_id and 'CS-' in str(case_id):
            cases.append({
                'row': row_idx,
                'id': str(case_id),
                'item': ws.cell(row_idx, 5).value,
                'precondition': ws.cell(row_idx, 6).value,
                'step': ws.cell(row_idx, 7).value,
                'expected': ws.cell(row_idx, 8).value,
            })
    return cases


def login_if_needed(page):
    """登录（如需要）"""
    if "login" in page.url.lower():
        page.get_by_title("账号登录").click()
        time.sleep(0.5)
        page.get_by_role("textbox", name="请输入手机号").fill(USERNAME)
        page.get_by_role("textbox", name="请输入密码").fill(PASSWORD)
        page.get_by_role("button", name="登 录").click()
        page.wait_for_load_state("domcontentloaded")
        time.sleep(3)


# ==================== 浏览器操作函数（按需修改） ====================
def open_create_modal(page):
    """打开创建弹窗"""
    # Step 1: 点击item激活
    page.evaluate("""
        () => {
            const d = document.querySelector('#myiframe').contentDocument;
            const items = d.querySelectorAll('.event-type');
            for (const item of items) {
                const label = item.querySelector('.event-type__label');
                if (label && label.textContent.includes('{目标事件类型}')) {
                    item.click();
                    return;
                }
            }
        }
    """)
    time.sleep(0.5)

    # Step 2: 获取action按钮坐标，用真实鼠标点击（JS .click() 不触发 Vue）
    result = page.evaluate("""
        () => {
            const d = document.querySelector('#myiframe').contentDocument;
            const items = d.querySelectorAll('.event-type');
            for (const item of items) {
                const label = item.querySelector('.event-type__label');
                if (label && label.textContent.includes('{目标事件类型}')) {
                    const btn = item.querySelector('.event-type__action');
                    if (!btn) return { found: false };
                    const r = btn.getBoundingClientRect();
                    const iframe = document.querySelector('#myiframe').getBoundingClientRect();
                    return {
                        found: true,
                        x: r.left + r.width / 2 + iframe.left,
                        y: r.top + r.height / 2 + iframe.top
                    };
                }
            }
            return { found: false };
        }
    """)
    if result.get("found"):
        page.mouse.click(result["x"], result["y"])
    time.sleep(2)


def reset_form(page):
    """重置表单到初始状态"""
    page.evaluate("""
        () => {
            const d = document.querySelector('#myiframe').contentDocument;
            const modal = d.querySelector('.mask-dialog-style-component');
            if (!modal) return;
            // 清空桩号
            const inputs = modal.querySelectorAll('input[role="spinbutton"]');
            for (const inp of inputs) {
                inp.value = '';
                inp.dispatchEvent(new Event('input', {bubbles: true}));
                inp.dispatchEvent(new Event('change', {bubbles: true}));
            }
            // 重置影响位置
            const locs = modal.querySelectorAll('.location-option, .location-item');
            for (const loc of locs) {
                loc.classList.remove('checked', 'active');
                const inp = loc.querySelector('input');
                if (inp) inp.checked = false;
            }
            // 重置占道信息
            const halves = modal.querySelectorAll('.event-road-half');
            for (const half of halves) {
                const openIcons = half.querySelectorAll('i.icon-close');
                for (const icon of openIcons) icon.click();
            }
        }
    """)
    time.sleep(0.5)


def fill_input(page, selector, value):
    """
    通用表单填写 - 用鼠标点击 + keyboard.type 确保 Vue 接收
    selector: CSS选择器
    """
    for attempt in range(3):
        result = page.evaluate(f"""
            () => {{
                const d = document.querySelector('#myiframe').contentDocument;
                const el = d.querySelector('{selector}');
                if (!el) return null;
                const r = el.getBoundingClientRect();
                const iframe = document.querySelector('#myiframe').getBoundingClientRect();
                return {{
                    x: r.left + r.width / 2 + iframe.left,
                    y: r.top + r.height / 2 + iframe.top
                }};
            }}
        """)
        if result:
            break
        time.sleep(0.5)

    if not result:
        return

    page.mouse.click(result["x"], result["y"])
    time.sleep(0.2)
    page.keyboard.type(str(value))
    time.sleep(0.3)


def get_form_state(page):
    """获取表单状态"""
    return page.evaluate("""
        () => {
            const d = document.querySelector('#myiframe').contentDocument;
            if (!d) return { modal: null, errors: [] };
            const modal = d.querySelector('.mask-dialog-style-component');
            const errs = d.querySelectorAll('.ant-form-item-explain-error');
            const errMsgs = [];
            for (const e of errs) errMsgs.push(e.textContent);
            return { modal: !!modal, errors: errMsgs };
        }
    """)


def click_confirm(page):
    """点击确定按钮"""
    result = page.evaluate("""
        () => {
            const d = document.querySelector('#myiframe').contentDocument;
            const btns = d.querySelectorAll('button');
            for (const btn of btns) {
                if (btn.textContent.trim().includes('确')) {
                    const r = btn.getBoundingClientRect();
                    return { found: true, x: r.left + r.width / 2, y: r.top + r.height / 2 };
                }
            }
            return { found: false };
        }
    """)
    if result.get("found"):
        page.mouse.click(result["x"], result["y"])
        time.sleep(2)
        return True
    return False


def switch_location(page, loc):
    """切换影响位置"""
    loc_map = {'主线': 0, '收费站': 1, '服务区': 2, '枢纽': 3}
    idx = loc_map.get(loc, 0)
    page.evaluate(f"""
        () => {{
            const d = document.querySelector('#myiframe').contentDocument;
            const options = d.querySelectorAll('.location-option, .location-item');
            if (options.length > {idx}) {{
                options[{idx}].click();
            }}
        }}
    """)
    time.sleep(0.5)


def fill_pile_number(page, km='800', hm='50'):
    """填写桩号 - 用 mouse.click + keyboard.type 确保 Vue 接收"""
    for attempt in range(3):
        result = page.evaluate("""
            () => {
                const d = document.querySelector('#myiframe').contentDocument;
                const modal = d.querySelector('.mask-dialog-style-component');
                if (!modal) return null;
                const spins = modal.querySelectorAll('input[role="spinbutton"]');
                if (spins.length < 2) return null;
                const iframe = document.querySelector('#myiframe').getBoundingClientRect();
                return {
                    kmX: spins[0].getBoundingClientRect().left + spins[0].getBoundingClientRect().width / 2 + iframe.left,
                    kmY: spins[0].getBoundingClientRect().top + spins[0].getBoundingClientRect().height / 2 + iframe.top,
                    hmX: spins[1].getBoundingClientRect().left + spins[1].getBoundingClientRect().width / 2 + iframe.left,
                    hmY: spins[1].getBoundingClientRect().top + spins[1].getBoundingClientRect().height / 2 + iframe.top,
                };
            }
        """)
        if result:
            break
        time.sleep(0.5)

    if not result:
        return

    page.mouse.click(result["kmX"], result["kmY"])
    time.sleep(0.2)
    page.keyboard.type(str(km))
    time.sleep(0.2)
    page.mouse.click(result["hmX"], result["hmY"])
    time.sleep(0.2)
    page.keyboard.type(str(hm))
    time.sleep(0.5)


def toggle_road_block(page, direction=0, lane=0):
    """切换占道信息"""
    page.evaluate(f"""
        () => {{
            const d = document.querySelector('#myiframe').contentDocument;
            const halves = d.querySelectorAll('.event-road-half');
            if (halves.length > {direction}) {{
                const icons = halves[{direction}].querySelectorAll('i.icon-open');
                if (icons.length > {lane}) {{
                    icons[{lane}].click();
                }}
            }}
        }}
    """)
    time.sleep(0.5)


# ==================== 用例执行（按需修改） ====================
def run_test(page, case_info):
    case_id = case_info['id']
    screenshot_path = os.path.join(BASE_SCREENSHOTS, safe_filename(case_id) + '.png')

    try:
        reset_form(page)

        if case_id == '{用例ID-001}':
            open_create_modal(page)
            page.screenshot(path=screenshot_path)
            state = get_form_state(page)
            if state.get('modal'):
                return 'PASS', '弹窗打开成功', screenshot_path
            return 'FAIL', '弹窗未打开', screenshot_path

        elif case_id == '{用例ID-002}':
            open_create_modal(page)
            time.sleep(1)
            # 关闭弹窗
            page.evaluate("""
                () => {
                    const d = document.querySelector('#myiframe').contentDocument;
                    const btns = d.querySelectorAll('button');
                    for (const btn of btns) {
                        if (btn.textContent.trim() === '取消') {
                            btn.click();
                            return;
                        }
                    }
                }
            """)
            time.sleep(1)
            state = get_form_state(page)
            if not state.get('modal'):
                return 'PASS', '弹窗已关闭', ''
            return 'FAIL', '弹窗未关闭', ''

        # ... 更多用例分支

        else:
            return 'SKIP', '未知用例', ''

    except Exception as e:
        return 'FAIL', str(e)[:50], ''


# ==================== 报告生成 ====================
def generate_report(cases, dst_excel):
    wb = load_workbook(dst_excel)
    ws = wb.active

    green = PatternFill(start_color="C6EFCE", fill_type="solid")
    red = PatternFill(start_color="FFC7CE", fill_type="solid")

    for case in cases:
        row = case['row']
        ws.cell(row, 9, case.get('result', 'N/A'))
        ws.cell(row, 10, case.get('remark', ''))

        screenshot = case.get('screenshot')
        if screenshot and os.path.exists(screenshot):
            try:
                img = XLImage(screenshot)
                img.width = 200
                img.height = 120
                img.anchor = OneCellAnchor()
                ws.add_image(img, f'K{row}')
                ws.row_dimensions[row].height = 80
            except Exception as e:
                print(f"[WARN] 截图嵌入失败: {e}")

        fill = green if case.get('result') == 'PASS' else red
        ws.cell(row, 9).fill = fill

    wb.save(dst_excel)
    print(f"Excel已保存: {dst_excel}")

    # HTML报告
    html_path = os.path.join(BASE_OUTPUT, 'test_report.html')
    pass_count = sum(1 for c in cases if c.get('result') == 'PASS')
    fail_count = sum(1 for c in cases if c.get('result') == 'FAIL')

    rows = ''
    for case in cases:
        cls = (case.get('result') or '').lower() or 'skip'
        screenshot = case.get('screenshot')
        img_tag = f'<img src="../screenshots/{os.path.basename(screenshot)}" class="screenshot-thumb" onclick="window.open(this.src)" />' if screenshot else ''

        rows += f'''<tr class="{cls}">
            <td>{case['id']}</td>
            <td>{case.get('item', '')}</td>
            <td>{str(case.get('step', ''))[:100]}</td>
            <td>{str(case.get('expected', ''))[:100]}</td>
            <td>{case.get('result', '')}</td>
            <td>{case.get('remark', '')}</td>
            <td>{img_tag}</td>
        </tr>'''

    html = f'''<!DOCTYPE html>
<html>
<head>
    <meta charset="utf-8">
    <title>测试报告</title>
    <style>
        body{{font-family:Arial;margin:20px;background:#f5f5f5}}
        table{{width:100%;border-collapse:collapse;background:#fff;margin-top:20px}}
        th,td{{padding:8px;border:1px solid #ddd;text-align:left}}
        th{{background:#4CAF50;color:white}}
        .pass{{background:#C6EFCE}}.fail{{background:#FFC7CE}}
        img{{max-width:200px;cursor:pointer}}
        .screenshot-thumb{{cursor:pointer;border:1px solid #ddd;margin:5px}}
    </style>
</head>
<body>
    <h1>{os.path.basename(dst_excel).replace('.xlsx', '')}</h1>
    <p>时间: {datetime.now().strftime("%Y-%m-%d %H:%M:%S")}</p>
    <p>用例总数: {len(cases)} | 通过: {pass_count} | 失败: {fail_count}</p>
    <table>
        <tr><th>用例ID</th><th>测试项</th><th>步骤</th><th>预期结果</th><th>执行结果</th><th>备注</th><th>截图</th></tr>
        {rows}
    </table>
    <script>document.querySelectorAll(".screenshot-thumb").forEach(img=>img.addEventListener("click",e=>window.open(e.target.src)))</script>
</body>
</html>'''

    with open(html_path, 'w', encoding='utf-8') as f:
        f.write(html)
    print(f"HTML报告: {html_path}")


# ==================== 主流程 ====================
def main():
    os.makedirs(BASE_OUTPUT, exist_ok=True)
    os.makedirs(BASE_SCREENSHOTS, exist_ok=True)

    # 复制源文件
    src = os.path.join(BASE_INPUT, CASE_FILE)
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    dst_excel = os.path.join(BASE_OUTPUT, f"结果_{timestamp}.xlsx")
    shutil.copy2(src, dst_excel)

    cases = load_cases()
    print(f"加载 {len(cases)} 条用例")

    p = sync_playwright().start()
    browser = p.chromium.connect_over_cdp(CDP_URL)
    page = browser.contexts[0].pages[0]

    try:
        page.goto(TARGET_URL, timeout=60000)
        page.wait_for_load_state("domcontentloaded")
        time.sleep(3)

        login_if_needed(page)
        print("开始执行测试...")

        for case in cases:
            print(f"[{case['id']}] {str(case['item'])[:30]}...")
            result, remark, screenshot = run_test(page, case)
            case['result'] = result
            case['remark'] = remark
            case['screenshot'] = screenshot
            print(f"  -> {result}: {remark}")

        print("\n=== 测试完成 ===")
        generate_report(cases, dst_excel)

    finally:
        browser.close()
        p.stop()


if __name__ == "__main__":
    main()
